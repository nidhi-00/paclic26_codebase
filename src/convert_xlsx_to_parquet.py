import argparse
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook
from tqdm import tqdm

from .common import file_sha256, write_json


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--sheet", default="DATA")
    parser.add_argument("--chunk-size", type=int, default=50000)
    parser.add_argument("--no-hash", action="store_true")
    return parser.parse_args()


def get_sheet(workbook, sheet_arg):
    if str(sheet_arg).isdigit():
        return workbook.worksheets[int(sheet_arg)]
    return workbook[str(sheet_arg)]


def clean_cell(x):
    """
    GECO Excel contains mixed types in some columns.
    To avoid PyArrow type crashes, store raw Excel cells as strings first.
    Later preprocessing scripts convert numeric columns back to numbers.
    """
    if x is None:
        return None
    if x == ".":
        return None
    return str(x)


def convert_xlsx_to_parquet(input_path, output_path, sheet_arg="DATA", chunk_size=50000, compute_hash=True):
    input_path = Path(input_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if output_path.exists():
        print(f"Removing old partial file: {output_path}")
        output_path.unlink()

    print(f"Opening Excel file: {input_path}")
    wb = load_workbook(input_path, read_only=True, data_only=True)
    ws = get_sheet(wb, sheet_arg)

    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    header = [str(h).strip() for h in header]

    writer = None
    buffer = []
    total_rows = 0
    chunk_id = 0

    for row in tqdm(rows, desc="Streaming Excel rows"):
        cleaned_row = [clean_cell(x) for x in row]
        buffer.append(cleaned_row)

        if len(buffer) >= chunk_size:
            df = pd.DataFrame(buffer, columns=header)

            # Force all columns to string-compatible Arrow type.
            table = pa.Table.from_pandas(
                df,
                schema=pa.schema([(col, pa.string()) for col in df.columns]),
                preserve_index=False,
            )

            if writer is None:
                writer = pq.ParquetWriter(output_path, table.schema)

            writer.write_table(table)
            total_rows += len(df)
            chunk_id += 1
            print(f"Wrote chunk {chunk_id}; total rows = {total_rows:,}")
            buffer = []

    if buffer:
        df = pd.DataFrame(buffer, columns=header)
        table = pa.Table.from_pandas(
            df,
            schema=pa.schema([(col, pa.string()) for col in df.columns]),
            preserve_index=False,
        )

        if writer is None:
            writer = pq.ParquetWriter(output_path, table.schema)

        writer.write_table(table)
        total_rows += len(df)
        chunk_id += 1
        print(f"Wrote final chunk {chunk_id}; total rows = {total_rows:,}")

    if writer is not None:
        writer.close()

    wb.close()

    metadata_path = output_path.with_suffix(".metadata.json")
    write_json(
        {
            "input": str(input_path),
            "output": str(output_path),
            "sheet": str(sheet_arg),
            "chunk_size": chunk_size,
            "rows": total_rows,
            "columns": header,
            "input_sha256": file_sha256(input_path) if compute_hash else None,
        },
        metadata_path,
    )
    print("Done.")
    print(f"Saved to: {output_path}")
    print(f"Total rows written: {total_rows:,}")
    print(f"Metadata: {metadata_path}")


def main():
    args = parse_args()
    convert_xlsx_to_parquet(
        input_path=args.input,
        output_path=args.output,
        sheet_arg=args.sheet,
        chunk_size=args.chunk_size,
        compute_hash=not args.no_hash,
    )


if __name__ == "__main__":
    main()
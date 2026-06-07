from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .common import read_table


def inspect_excel(path: Path, sheet: str | int | None, rows: int) -> None:
    wb = load_workbook(filename=path, read_only=True, data_only=True)
    print("Sheets:", wb.sheetnames)
    ws = wb[wb.sheetnames[int(sheet)]] if isinstance(sheet, int) else wb[sheet or wb.sheetnames[0]]
    iterator = ws.iter_rows(values_only=True)
    header = next(iterator)
    print("\nColumns:")
    for i, col in enumerate(header):
        print(f"  {i:03d}: {col}")
    print(f"\nFirst {rows} rows:")
    for idx, row in zip(range(rows), iterator):
        print(dict(zip(header, row)))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--sheet", default=0)
    parser.add_argument("--rows", type=int, default=5)
    args = parser.parse_args()

    path = Path(args.input)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            sheet = int(args.sheet)
        except ValueError:
            sheet = args.sheet
        inspect_excel(path, sheet, args.rows)
    else:
        df = read_table(path, nrows=args.rows)
        print("Columns:")
        for i, col in enumerate(df.columns):
            print(f"  {i:03d}: {col}")
        print(f"\nFirst {args.rows} rows:")
        print(df.head(args.rows).to_string())


if __name__ == "__main__":
    main()
